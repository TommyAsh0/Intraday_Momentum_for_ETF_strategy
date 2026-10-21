#!/usr/bin/env python
"""把本工程 reports/ 下的叙述性 md 报告（含图、表、粗体、引用）转换为 docx。

用法：
    python scripts/md_to_docx.py reports/研究之路_旗舰策略.md
    python scripts/md_to_docx.py reports/研究之路_旗舰策略.md reports/研究之路_券商改良策略接入.md
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
IMG_RE = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")
ROW_RE = re.compile(r"^\|(.+)\|\s*$")


def add_runs(p: docx.text.paragraph.Paragraph, text: str) -> None:
    """极简行内格式：**bold** 与 `code`。"""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = p.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = p.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        else:
            p.add_run(tok)


def md_to_docx(md_path: Path) -> Path:
    doc = docx.Document()
    # 基础样式
    st = doc.styles["Normal"]
    st.font.name = "Microsoft YaHei"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(docx.oxml.ns.qn("w:eastAsia"), "Microsoft YaHei")

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1
            continue
        # 标题
        if s.startswith("#### "):
            doc.add_heading(s[5:], level=4)
        elif s.startswith("### "):
            doc.add_heading(s[4:], level=3)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("# "):
            h = doc.add_heading("", level=0)
            add_runs(h, s[2:])
        # 图片
        elif (m := IMG_RE.match(s)):
            alt, src = m.group(1), m.group(2)
            img = (md_path.parent / src).resolve()
            if img.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(img), width=Cm(15.5))
                if alt:
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = cap.add_run(alt)
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
        # 表格
        elif ROW_RE.match(s):
            rows = []
            while i < len(lines) and ROW_RE.match(lines[i].strip()):
                cells = [c.strip() for c in ROW_RE.match(lines[i].strip()).group(1).split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                t.style = "Light Grid Accent 1"
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        t.cell(ri, ci).text = cell
                        for par in t.cell(ri, ci).paragraphs:
                            for run in par.runs:
                                run.font.size = Pt(9)
                                if ri == 0:
                                    run.bold = True
            continue
        # 分隔线
        elif re.fullmatch(r"-{3,}", s):
            doc.add_paragraph("─" * 40).alignment = WD_ALIGN_PARAGRAPH.CENTER
        # 引用块
        elif s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            r = p.add_run(s.lstrip("> ").strip())
            r.italic = True
            r.font.color.rgb = RGBColor(0x50, 0x50, 0x50)
        # 无序列表
        elif s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, s[2:])
        elif re.match(r"^\d+\.\s", s):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s", "", s))
        else:
            p = doc.add_paragraph()
            add_runs(p, s)
        i += 1

    out = md_path.with_suffix(".docx")
    doc.save(out)
    return out


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    targets = [Path(a) for a in sys.argv[1:]] or [
        ROOT / "reports" / "研究之路_旗舰策略.md",
        ROOT / "reports" / "研究之路_券商改良策略接入.md",
    ]
    for md in targets:
        md = md if md.is_absolute() else (ROOT / md)
        out = md_to_docx(md)
        print(f"OK: {md.name} -> {out.name} ({out.stat().st_size // 1024} KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
